#!/usr/bin/env python3
"""
Grand recon report (DuckDB extracts -> PeopleSoft Oracle validation).

What this does:
- Reads your extracted DuckDB tables:
  - po_header
  - goods_po_line
  - service_po_line
- Builds the PO list from DuckDB (union of all POs in those tables)
- Queries PeopleSoft (Oracle) for the same PO list and recomputes:
  - PO header status/vendor/date
  - Service vs goods flag (account -> PO type mapping)
  - Total PO amounts/qty, received amounts/qty
  - Vouchered (matched + PAID) amounts/qty
  - Remaining amounts/qty using the same gating logic you’ve been using
- Produces an Excel report with:
  - Summary
  - DuckDB totals by PO
  - Oracle totals by PO
  - Joined comparison + mismatch flags
  - Missing POs (DuckDB-only / Oracle-only)

Usage (PowerShell):
  python bh_oracle_extracts\\duckdb_oracle_grand_recon_report.py `
    --duckdb "C:\\PT8.61.09_Client_ORA\\python\\PO\\extracts.duckdb" `
    --oracle-user "vlombard" `
    --oracle-password "Playp1ace" `
    --oracle-connect "vms-00-00-773.bhsi.com:1521/ERPWD1" `
    --output "C:\\PT8.61.09_Client_ORA\\python\\PO\\grand_recon_from_duckdb.xlsx" `
    --asof-date 2026-01-15
"""

from __future__ import annotations

import argparse
import csv
import os
import time
from dataclasses import dataclass
from datetime import datetime
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import duckdb

try:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Alignment, Font  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    _HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    Workbook = None  # type: ignore
    Alignment = None  # type: ignore
    Font = None  # type: ignore
    get_column_letter = None  # type: ignore
    _HAVE_OPENPYXL = False

try:
    import oracledb  # type: ignore
except Exception:  # pragma: no cover
    oracledb = None  # type: ignore


def _write_csv(path: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        if headers:
            w.writerow(list(headers))
        for r in rows:
            w.writerow(list(r))


def _autosize(ws, max_width: int = 60) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), max_width)


def _write_table(ws, start_row: int, start_col: int, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    header_font = Font(bold=True)
    for j, h in enumerate(headers, start_col):
        cell = ws.cell(row=start_row, column=j, value=h)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, start_col):
            ws.cell(row=i, column=j, value=v)

    _autosize(ws, max_width=80)


def _oracle_connect(user: str, password: str, connect_str: str):
    if oracledb is None:
        raise RuntimeError("Missing dependency: oracledb. Install with: pip install oracledb")
    return oracledb.connect(user=user, password=password, dsn=connect_str)


def _duckdb_exists_table(con: duckdb.DuckDBPyConnection, name: str) -> bool:
    rows = con.execute("SELECT COUNT(*) FROM information_schema.tables WHERE table_name = ?", [name]).fetchone()
    return bool(rows and rows[0] > 0)


def _duckdb_num(expr: str) -> str:
    # DuckDB tables are often all VARCHAR; use try_cast safely.
    return f"COALESCE(try_cast(replace(trim({expr}), ',', '') AS DOUBLE), 0.0)"


@dataclass
class DuckPO:
    po_id: str
    goods_line_cnt: int
    service_line_cnt: int
    goods_amt: float
    service_amt: float
    goods_qty: float


def _load_duckdb_po_totals(con: duckdb.DuckDBPyConnection) -> Tuple[List[DuckPO], List[str]]:
    required = ["po_header", "goods_po_line", "service_po_line"]
    for t in required:
        if not _duckdb_exists_table(con, t):
            raise RuntimeError(f"Missing DuckDB table: {t!r}. Expected tables: {', '.join(required)}")

    # Union PO ids from all three tables (header + lines)
    po_list_rows = con.execute(
        """
        WITH u AS (
          SELECT DISTINCT no AS po_id FROM po_header
          UNION
          SELECT DISTINCT no AS po_id FROM goods_po_line
          UNION
          SELECT DISTINCT no AS po_id FROM service_po_line
        )
        SELECT po_id
        FROM u
        WHERE po_id IS NOT NULL AND trim(po_id) <> ''
        ORDER BY po_id
        """
    ).fetchall()
    po_list = [str(r[0]) for r in po_list_rows if r and r[0] is not None]

    rows = con.execute(
        f"""
        WITH
        g AS (
          SELECT
            no AS po_id,
            COUNT(*)::BIGINT AS goods_line_cnt,
            SUM({_duckdb_num('extended_amount')}) AS goods_amt,
            SUM({_duckdb_num('quantity')}) AS goods_qty
          FROM goods_po_line
          GROUP BY no
        ),
        s AS (
          SELECT
            no AS po_id,
            COUNT(*)::BIGINT AS service_line_cnt,
            SUM({_duckdb_num('extended_amount')}) AS service_amt
          FROM service_po_line
          GROUP BY no
        ),
        u AS (
          SELECT po_id FROM g
          UNION
          SELECT po_id FROM s
          UNION
          SELECT DISTINCT no AS po_id FROM po_header
        )
        SELECT
          u.po_id,
          COALESCE(g.goods_line_cnt, 0) AS goods_line_cnt,
          COALESCE(s.service_line_cnt, 0) AS service_line_cnt,
          COALESCE(g.goods_amt, 0.0) AS goods_amt,
          COALESCE(s.service_amt, 0.0) AS service_amt,
          COALESCE(g.goods_qty, 0.0) AS goods_qty
        FROM u
        LEFT JOIN g ON g.po_id = u.po_id
        LEFT JOIN s ON s.po_id = u.po_id
        WHERE u.po_id IS NOT NULL AND trim(u.po_id) <> ''
        ORDER BY u.po_id
        """
    ).fetchall()

    out: List[DuckPO] = []
    for r in rows:
        out.append(
            DuckPO(
                po_id=str(r[0]),
                goods_line_cnt=int(r[1] or 0),
                service_line_cnt=int(r[2] or 0),
                goods_amt=float(r[3] or 0.0),
                service_amt=float(r[4] or 0.0),
                goods_qty=float(r[5] or 0.0),
            )
        )
    return out, po_list


def _oracle_fetch_po_recon_rows(
    conn,
    po_list: Sequence[str],
    asof_date: Optional[str],
    chunk_size: int,
) -> Tuple[List[str], List[Sequence[object]]]:
    """
    Returns one row per PO with:
    - has_service flag
    - PO header fields
    - totals: qty_po/amt_po, qty_rcvd/amt_rcvd, qty_vchr/amt_vchr (matched+paid)
    - remaining amounts/qty per rules
    - reason included
    """

    oracle_sql = """
WITH
params AS (
  SELECT TRUNC(SYSDATE) AS asof_dt,
         ADD_MONTHS(TRUNC(SYSDATE), -12) AS lookback_dt
  FROM dual
),
duckdb_po_list AS (
  SELECT column_value AS po_id
  FROM TABLE(:po_list)
),
hdr_candidates AS (
  SELECT /*+ MATERIALIZE */
         h.business_unit,
         h.po_id,
         h.po_dt,
         h.po_status,
         h.vendor_id
  FROM ps_po_hdr h
  JOIN duckdb_po_list d
    ON d.po_id = h.po_id
  CROSS JOIN params p
  WHERE h.po_dt <= p.asof_dt
    AND h.po_dt >= p.lookback_dt
    AND h.po_status NOT IN ('C','X')
    AND h.vendor_id <> '2000017041'
),
service_flags AS (
  SELECT
      d.business_unit,
      d.po_id,
      CASE WHEN MAX(CASE WHEN x.bh_xwlk_t1 IS NOT NULL THEN 1 ELSE 0 END) = 1 THEN 'Y' ELSE 'N' END AS has_service
  FROM hdr_candidates hc
  JOIN ps_po_line_distrib d
    ON d.business_unit = hc.business_unit
   AND d.po_id         = hc.po_id
  LEFT JOIN ps_bh_xwlk_val_tbl x
    ON x.longname       = 'WD_ACCT_TO_PO_TYPE'
   AND x.bh_xwlk_module = 'PO'
   AND x.bh_xwlk_track  = 'SCM'
   AND x.bh_xwlk_s2     = d.account
  GROUP BY d.business_unit, d.po_id
),
paid_vouchers AS (
  SELECT DISTINCT px.business_unit, px.voucher_id
  FROM ps_pymnt_vchr_xref px
  JOIN ps_payment_tbl pt
    ON pt.bank_setid    = px.bank_setid
   AND pt.bank_cd       = px.bank_cd
   AND pt.bank_acct_key = px.bank_acct_key
   AND pt.pymnt_id      = px.pymnt_id
   AND pt.schedule_id   = px.schedule_id
  CROSS JOIN params p
  WHERE px.pymnt_action <> 'X'
    AND pt.pymnt_status = 'P'
    AND pt.pymnt_dt < (p.asof_dt + 1)
),
vchr_sum_match AS (
  SELECT
      vl.business_unit_po AS business_unit,
      vl.po_id,
      vl.line_nbr,
      NVL(vl.sched_nbr, 1) AS sched_nbr,
      SUM(NVL(vl.merchandise_amt, 0)) AS merch_amt_vchr,
      SUM(NVL(vl.qty_vchr, 0))        AS qty_vchr
  FROM ps_voucher_line vl
  JOIN ps_voucher v
    ON v.business_unit = vl.business_unit
   AND v.voucher_id    = vl.voucher_id
  JOIN paid_vouchers pv
    ON pv.business_unit = v.business_unit
   AND pv.voucher_id    = v.voucher_id
  JOIN hdr_candidates hc
    ON hc.business_unit = vl.business_unit_po
   AND hc.po_id         = vl.po_id
  CROSS JOIN params p
  WHERE v.entry_status <> 'X'
    AND NVL(v.match_status_vchr, ' ') = 'M'
    AND vl.po_id IS NOT NULL
    AND vl.po_id <> ' '
    AND NVL(v.invoice_dt, v.entered_dt) < (p.asof_dt + 1)
  GROUP BY vl.business_unit_po, vl.po_id, vl.line_nbr, NVL(vl.sched_nbr, 1)
),
recv_agg AS (
  SELECT
      r.business_unit_po AS business_unit,
      r.po_id,
      r.line_nbr,
      r.sched_nbr,
      SUM(NVL(r.qty_sh_accpt, 0)) AS qty_rcvd,
      SUM(NVL(r.merchandise_amt_po, 0)) AS amt_rcvd
  FROM ps_recv_ln_ship r
  JOIN hdr_candidates hc
    ON hc.business_unit = r.business_unit_po
   AND hc.po_id         = r.po_id
  CROSS JOIN params p
  WHERE r.recv_ship_status <> 'X'
    AND r.receipt_dttm < CAST(p.asof_dt + 1 AS TIMESTAMP)
  GROUP BY r.business_unit_po, r.po_id, r.line_nbr, r.sched_nbr
),
po_line_flags AS (
  SELECT
      l.business_unit,
      l.po_id,
      l.line_nbr,
      NVL(l.recv_req, 'Y')     AS recv_req,
      NVL(l.amt_only_flg, 'N') AS amt_only_flg
  FROM ps_po_line l
  JOIN hdr_candidates hc
    ON hc.business_unit = l.business_unit
   AND hc.po_id         = l.po_id
  WHERE l.cancel_status <> 'X'
),
sched_facts AS (
  SELECT
      s.business_unit,
      s.po_id,
      s.line_nbr,
      s.sched_nbr,
      lf.recv_req,
      lf.amt_only_flg,
      NVL(s.qty_po, 0) AS qty_po,
      NVL(s.price_po, 0) AS price_po,
      NVL(s.merchandise_amt, NVL(s.qty_po,0) * NVL(s.price_po,0)) AS sched_amt,
      NVL(r.qty_rcvd, 0) AS qty_rcvd,
      NVL(r.amt_rcvd, 0) AS amt_rcvd,
      NVL(vs.qty_vchr, 0) AS qty_vchr,
      NVL(vs.merch_amt_vchr, 0) AS amt_vchr,
      GREATEST(NVL(s.merchandise_amt, NVL(s.qty_po,0) * NVL(s.price_po,0)) - NVL(vs.merch_amt_vchr,0), 0) AS rem_amt,
      GREATEST(NVL(s.qty_po,0) - NVL(vs.qty_vchr,0), 0) AS rem_qty
  FROM ps_po_line_ship s
  JOIN hdr_candidates hc
    ON hc.business_unit = s.business_unit
   AND hc.po_id         = s.po_id
  JOIN po_line_flags lf
    ON lf.business_unit = s.business_unit
   AND lf.po_id         = s.po_id
   AND lf.line_nbr      = s.line_nbr
  JOIN ps_po_line_distrib d
    ON d.business_unit      = s.business_unit
   AND d.po_id              = s.po_id
   AND d.line_nbr           = s.line_nbr
   AND d.sched_nbr          = s.sched_nbr
   AND d.distrib_ln_status <> 'X'
  LEFT JOIN recv_agg r
    ON r.business_unit = s.business_unit
   AND r.po_id         = s.po_id
   AND r.line_nbr      = s.line_nbr
   AND r.sched_nbr     = s.sched_nbr
  LEFT JOIN vchr_sum_match vs
    ON vs.business_unit = s.business_unit
   AND vs.po_id         = s.po_id
   AND vs.line_nbr      = s.line_nbr
   AND vs.sched_nbr     = s.sched_nbr
  WHERE NVL(s.cancel_status,' ') NOT IN ('C','X')
),
po_rollup AS (
  SELECT
      sf.business_unit,
      sf.po_id,
      MAX(svc.has_service) AS has_service,
      SUM(sf.qty_po) AS qty_po_total,
      SUM(sf.sched_amt) AS amt_po_total,
      SUM(sf.qty_rcvd) AS qty_rcvd_total,
      SUM(sf.amt_rcvd) AS amt_rcvd_total,
      SUM(sf.qty_vchr) AS qty_vchr_paid_matched_total,
      SUM(sf.amt_vchr) AS amt_vchr_paid_matched_total,
      SUM(CASE WHEN svc.has_service = 'Y' THEN sf.rem_amt
               WHEN sf.amt_only_flg = 'Y' THEN sf.rem_amt
               ELSE sf.rem_qty * sf.price_po
          END) AS total_remaining_amt,
      SUM(CASE WHEN svc.has_service = 'N' AND sf.amt_only_flg <> 'Y' THEN sf.rem_qty ELSE 0 END) AS goods_remaining_qty,
      MAX(CASE WHEN svc.has_service = 'Y' AND sf.rem_amt > 0 THEN 1 ELSE 0 END) AS has_service_open,
      MAX(CASE WHEN svc.has_service = 'N' AND sf.amt_only_flg = 'Y' AND sf.rem_amt > 0 THEN 1 ELSE 0 END) AS has_goods_amt_open,
      MAX(CASE WHEN svc.has_service = 'N' AND sf.amt_only_flg <> 'Y' AND sf.rem_qty > 0 THEN 1 ELSE 0 END) AS has_goods_qty_open
  FROM sched_facts sf
  JOIN service_flags svc
    ON svc.business_unit = sf.business_unit
   AND svc.po_id         = sf.po_id
  GROUP BY sf.business_unit, sf.po_id
)
SELECT
  hc.business_unit,
  hc.po_id,
  hc.po_dt,
  hc.po_status,
  hc.vendor_id,
  pr.has_service,
  pr.qty_po_total,
  pr.amt_po_total,
  pr.qty_rcvd_total,
  pr.amt_rcvd_total,
  pr.qty_vchr_paid_matched_total,
  pr.amt_vchr_paid_matched_total,
  pr.goods_remaining_qty,
  pr.total_remaining_amt,
  CASE
    WHEN pr.has_service_open = 1 THEN 'SERVICE_REMAINING_AMT_GT_0'
    WHEN pr.has_goods_amt_open = 1 THEN 'GOODS_AMT_ONLY_REMAINING_AMT_GT_0'
    WHEN pr.has_goods_qty_open = 1 THEN 'GOODS_QTY_REMAINING_GT_0'
    ELSE 'NOT_OPEN_BY_RULES'
  END AS oracle_reason_included
FROM hdr_candidates hc
JOIN po_rollup pr
  ON pr.business_unit = hc.business_unit
 AND pr.po_id         = hc.po_id
ORDER BY pr.total_remaining_amt DESC, hc.business_unit, hc.po_id
"""

    if asof_date:
        oracle_sql = oracle_sql.replace(
            "SELECT TRUNC(SYSDATE) AS asof_dt,",
            "SELECT TRUNC(TO_DATE(:asof_dt, 'YYYY-MM-DD')) AS asof_dt,",
        ).replace(
            "ADD_MONTHS(TRUNC(SYSDATE), -12) AS lookback_dt",
            "ADD_MONTHS(TRUNC(TO_DATE(:asof_dt, 'YYYY-MM-DD')), -12) AS lookback_dt",
        )

    cur = conn.cursor()
    all_rows: List[Sequence[object]] = []
    cols: List[str] = []
    try:
        for i in range(0, len(po_list), chunk_size):
            chunk = list(po_list[i : i + chunk_size])
            po_var = cur.arrayvar(oracledb.DB_TYPE_VARCHAR, chunk, type_name="SYS.ODCIVARCHAR2LIST")
            bind = {"po_list": po_var}
            if asof_date:
                bind["asof_dt"] = asof_date
            cur.execute(oracle_sql, bind)
            if not cols and cur.description:
                cols = [d[0] for d in cur.description]
            all_rows.extend(cur.fetchall())
    finally:
        cur.close()
    return cols, all_rows


def _rows_to_dict(rows: Sequence[Sequence[object]], key_idx: int) -> Dict[str, Sequence[object]]:
    out: Dict[str, Sequence[object]] = {}
    for r in rows:
        if not r or r[key_idx] is None:
            continue
        out[str(r[key_idx])] = r
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="Grand recon: DuckDB extracts vs PeopleSoft Oracle validation")
    ap.add_argument("--duckdb", required=True, help="Path to DuckDB database file (.duckdb)")
    ap.add_argument("--duckdb-read-only", action="store_true", help="Open DuckDB in read-only mode")
    ap.add_argument("--oracle-user", required=True)
    ap.add_argument("--oracle-password", required=True)
    ap.add_argument("--oracle-connect", required=True, help="host:port/service_name")
    ap.add_argument("--asof-date", default=None, help="YYYY-MM-DD (optional; if omitted Oracle uses SYSDATE)")
    ap.add_argument("--oracle-chunk-size", type=int, default=2000, help="POs per Oracle query chunk (default: 2000)")
    ap.add_argument("--output", required=True, help="Output .xlsx path (or a directory for CSV fallback)")
    ap.add_argument("--max-rows", type=int, default=200000, help="Max rows written per sheet (default: 200k)")
    ap.add_argument("--amt-tol", type=float, default=0.01, help="Amount tolerance for mismatches (default: 0.01)")
    ap.add_argument("--qty-tol", type=float, default=0.0001, help="Qty tolerance for mismatches (default: 0.0001)")
    args = ap.parse_args()

    duckdb_path = os.path.abspath(args.duckdb)
    out_path = os.path.abspath(args.output)

    # ---- DuckDB ----
    dcon = duckdb.connect(duckdb_path, read_only=args.duckdb_read_only)
    try:
        t0 = time.time()
        duck_pos, po_list = _load_duckdb_po_totals(dcon)
        duck_ms = int((time.time() - t0) * 1000)
    finally:
        dcon.close()

    # ---- Oracle ----
    ocon = _oracle_connect(args.oracle_user, args.oracle_password, args.oracle_connect)
    try:
        t1 = time.time()
        ora_cols, ora_rows = _oracle_fetch_po_recon_rows(ocon, po_list, args.asof_date, args.oracle_chunk_size)
        ora_ms = int((time.time() - t1) * 1000)
    finally:
        try:
            ocon.close()
        except Exception:
            pass

    # ---- Join ----
    duck_by_po: Dict[str, DuckPO] = {d.po_id: d for d in duck_pos}
    if not ora_cols:
        raise SystemExit("Oracle query returned no columns (unexpected).")
    ora_po_idx = [c.lower() for c in ora_cols].index("po_id") if "po_id" in [c.lower() for c in ora_cols] else 1
    ora_by_po = _rows_to_dict(ora_rows, ora_po_idx)

    joined_headers = [
        "po_id",
        "duckdb_goods_amt",
        "duckdb_service_amt",
        "duckdb_total_amt",
        "duckdb_goods_qty",
        "oracle_total_remaining_amt",
        "oracle_goods_remaining_qty",
        "oracle_has_service",
        "oracle_reason_included",
        "oracle_po_status",
        "oracle_po_dt",
        "oracle_vendor_id",
        "amt_diff",
        "qty_diff",
        "amt_match",
        "qty_match",
    ]

    joined_rows: List[List[object]] = []
    mismatch_rows: List[List[object]] = []
    missing_in_oracle: List[List[object]] = []
    missing_in_duckdb: List[List[object]] = []

    # map oracle columns we care about (case-insensitive)
    lc = [c.lower() for c in ora_cols]
    def _idx(name: str) -> int:
        return lc.index(name.lower())

    idx_has_service = _idx("has_service")
    idx_reason = _idx("oracle_reason_included")
    idx_po_status = _idx("po_status")
    idx_po_dt = _idx("po_dt")
    idx_vendor = _idx("vendor_id")
    idx_rem_amt = _idx("total_remaining_amt")
    idx_rem_qty = _idx("goods_remaining_qty")

    all_po_ids = sorted(set(list(duck_by_po.keys()) + list(ora_by_po.keys())))
    for po_id in all_po_ids:
        d = duck_by_po.get(po_id)
        o = ora_by_po.get(po_id)

        if d is None and o is not None:
            missing_in_duckdb.append([po_id, "Present in Oracle result but missing from DuckDB extracts"])
            continue
        if d is not None and o is None:
            missing_in_oracle.append([po_id, d.goods_amt + d.service_amt, d.goods_line_cnt, d.service_line_cnt])
            continue

        assert d is not None and o is not None
        duck_total = float(d.goods_amt + d.service_amt)
        ora_total = float(o[idx_rem_amt] or 0.0)
        duck_qty = float(d.goods_qty or 0.0)
        ora_qty = float(o[idx_rem_qty] or 0.0)

        amt_diff = duck_total - ora_total
        qty_diff = duck_qty - ora_qty
        amt_match = abs(amt_diff) <= float(args.amt_tol)
        qty_match = abs(qty_diff) <= float(args.qty_tol)

        row = [
            po_id,
            d.goods_amt,
            d.service_amt,
            duck_total,
            duck_qty,
            ora_total,
            ora_qty,
            o[idx_has_service],
            o[idx_reason],
            o[idx_po_status],
            o[idx_po_dt],
            o[idx_vendor],
            amt_diff,
            qty_diff,
            "PASS" if amt_match else "FAIL",
            "PASS" if qty_match else "FAIL",
        ]
        joined_rows.append(row)
        if not amt_match or (duck_qty > 0 and not qty_match):
            mismatch_rows.append(row)

    # ---- Output ----
    if out_path.lower().endswith(".xlsx"):
        if not _HAVE_OPENPYXL:
            raise SystemExit("openpyxl not installed. Install with: pip install openpyxl (or output to a directory for CSV fallback)")

        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)

        summary = wb.create_sheet("Summary")
        summary.append(["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        summary.append(["DuckDB", duckdb_path])
        summary.append(["Oracle connect", args.oracle_connect])
        summary.append(["As-of date", args.asof_date or "SYSDATE"])
        summary.append([])
        summary.append(["DuckDB POs (union)", len(po_list)])
        summary.append(["DuckDB load ms", duck_ms])
        summary.append(["Oracle rows returned", len(ora_rows)])
        summary.append(["Oracle query ms", ora_ms])
        summary.append(["Joined rows", len(joined_rows)])
        summary.append(["Mismatches (amt or qty)", len(mismatch_rows)])
        summary.append(["Missing in Oracle", len(missing_in_oracle)])
        summary.append(["Missing in DuckDB", len(missing_in_duckdb)])
        _autosize(summary, max_width=80)

        ws_duck = wb.create_sheet("DuckDB_PO_Totals")
        duck_headers = ["po_id", "goods_line_cnt", "service_line_cnt", "goods_amt", "service_amt", "goods_qty", "duckdb_total_amt"]
        duck_rows = [
            [d.po_id, d.goods_line_cnt, d.service_line_cnt, d.goods_amt, d.service_amt, d.goods_qty, d.goods_amt + d.service_amt]
            for d in duck_pos[: args.max_rows]
        ]
        _write_table(ws_duck, 1, 1, duck_headers, duck_rows)

        ws_or = wb.create_sheet("Oracle_PO_Totals")
        _write_table(ws_or, 1, 1, ora_cols, ora_rows[: args.max_rows])

        ws_join = wb.create_sheet("Compare")
        _write_table(ws_join, 1, 1, joined_headers, joined_rows[: args.max_rows])

        ws_mm = wb.create_sheet("Mismatches")
        _write_table(ws_mm, 1, 1, joined_headers, mismatch_rows[: args.max_rows])

        ws_mo = wb.create_sheet("Missing_in_Oracle")
        _write_table(ws_mo, 1, 1, ["po_id", "duckdb_total_amt", "duckdb_goods_line_cnt", "duckdb_service_line_cnt"], missing_in_oracle[: args.max_rows])

        ws_md = wb.create_sheet("Missing_in_DuckDB")
        _write_table(ws_md, 1, 1, ["po_id", "note"], missing_in_duckdb[: args.max_rows])

        wb.save(out_path)
        print(f"✓ Wrote Excel report: {out_path}")
        return 0

    # CSV fallback: output treated as a directory
    out_dir = out_path
    os.makedirs(out_dir, exist_ok=True)
    _write_csv(
        os.path.join(out_dir, "summary.csv"),
        ["key", "value"],
        [
            ("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("duckdb", duckdb_path),
            ("oracle_connect", args.oracle_connect),
            ("asof_date", args.asof_date or "SYSDATE"),
            ("duckdb_po_cnt", len(po_list)),
            ("oracle_rows", len(ora_rows)),
            ("joined_rows", len(joined_rows)),
            ("mismatches", len(mismatch_rows)),
            ("missing_in_oracle", len(missing_in_oracle)),
            ("missing_in_duckdb", len(missing_in_duckdb)),
        ],
    )
    _write_csv(os.path.join(out_dir, "duckdb_po_totals.csv"), ["po_id", "goods_line_cnt", "service_line_cnt", "goods_amt", "service_amt", "goods_qty", "duckdb_total_amt"], duck_rows)
    _write_csv(os.path.join(out_dir, "oracle_po_totals.csv"), ora_cols, ora_rows[: args.max_rows])
    _write_csv(os.path.join(out_dir, "compare.csv"), joined_headers, joined_rows[: args.max_rows])
    _write_csv(os.path.join(out_dir, "mismatches.csv"), joined_headers, mismatch_rows[: args.max_rows])
    _write_csv(os.path.join(out_dir, "missing_in_oracle.csv"), ["po_id", "duckdb_total_amt", "duckdb_goods_line_cnt", "duckdb_service_line_cnt"], missing_in_oracle[: args.max_rows])
    _write_csv(os.path.join(out_dir, "missing_in_duckdb.csv"), ["po_id", "note"], missing_in_duckdb[: args.max_rows])
    print(f"✓ Wrote CSV outputs to: {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

