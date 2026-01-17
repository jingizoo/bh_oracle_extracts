#!/usr/bin/env python3
"""
Run the full "grand recon" end-to-end:
- Execute DuckDB statements from `grand_recon.sql`
- Use the resulting PO list to execute the Oracle validation query (same file)
- Write a single Excel report (counts + sums + per-PO reasons)

Usage (PowerShell):
  python bh_oracle_extracts\\run_grand_recon.py `
    --duckdb "C:\\path\\extracts.duckdb" `
    --oracle-user "vlombard" `
    --oracle-password "Playp1ace" `
    --oracle-connect "vms-00-00-773.bhsi.com:1521/ERPWD1" `
    --output "C:\\path\\grand_recon.xlsx" `
    --asof-date 2026-01-15
"""

from __future__ import annotations

import argparse
import os
import re
import time
from dataclasses import dataclass
from datetime import datetime
import csv
from typing import Iterable, List, Optional, Sequence, Tuple

import duckdb

try:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Alignment, Font  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    _HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    Workbook = None  # type: ignore
    Alignment = None  # type: ignore
    Font = None  # type: ignore
    get_column_letter = None  # type: ignore
    _HAVE_OPENPYXL = False


try:
    import oracledb  # type: ignore
except Exception:  # pragma: no cover
    oracledb = None  # type: ignore


def split_sql_statements(sql_text: str) -> List[str]:
    """
    Split SQL text into statements on semicolons, ignoring semicolons inside:
    - single-quoted strings
    - double-quoted identifiers
    - line comments (-- ...)
    - block comments (/* ... */)
    """
    statements: List[str] = []
    buf: List[str] = []

    in_single = False
    in_double = False
    in_line_comment = False
    in_block_comment = False

    i = 0
    n = len(sql_text)
    while i < n:
        ch = sql_text[i]
        nxt = sql_text[i + 1] if i + 1 < n else ""

        if in_line_comment:
            buf.append(ch)
            if ch == "\n":
                in_line_comment = False
            i += 1
            continue

        if in_block_comment:
            buf.append(ch)
            if ch == "*" and nxt == "/":
                buf.append(nxt)
                i += 2
                in_block_comment = False
            else:
                i += 1
            continue

        if not in_single and not in_double:
            if ch == "-" and nxt == "-":
                buf.append(ch)
                buf.append(nxt)
                i += 2
                in_line_comment = True
                continue
            if ch == "/" and nxt == "*":
                buf.append(ch)
                buf.append(nxt)
                i += 2
                in_block_comment = True
                continue

        if ch == "'" and not in_double:
            buf.append(ch)
            if in_single and nxt == "'":
                buf.append(nxt)
                i += 2
                continue
            in_single = not in_single
            i += 1
            continue

        if ch == '"' and not in_single:
            buf.append(ch)
            if in_double and nxt == '"':
                buf.append(nxt)
                i += 2
                continue
            in_double = not in_double
            i += 1
            continue

        if ch == ";" and not in_single and not in_double:
            stmt = "".join(buf).strip()
            if stmt:
                statements.append(stmt)
            buf = []
            i += 1
            continue

        buf.append(ch)
        i += 1

    tail = "".join(buf).strip()
    if tail:
        statements.append(tail)
    return statements


def _write_table(ws, start_row: int, start_col: int, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    header_font = Font(bold=True)
    for j, h in enumerate(headers, start_col):
        cell = ws.cell(row=start_row, column=j, value=h)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, start_col):
            ws.cell(row=i, column=j, value=v)

    for j in range(start_col, start_col + len(headers)):
        max_len = 0
        col_letter = get_column_letter(j)
        for cell in ws[col_letter]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(10, max_len + 2), 60)


def _write_csv(path: str, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        if headers:
            w.writerow(list(headers))
        for r in rows:
            w.writerow(list(r))


def _extract_section(full_sql: str, start_marker: str, end_marker: str) -> str:
    s = full_sql
    a = s.find(start_marker)
    if a < 0:
        raise ValueError(f"Start marker not found: {start_marker!r}")
    b = s.find(end_marker, a)
    if b < 0:
        raise ValueError(f"End marker not found: {end_marker!r}")
    return s[a:b]


def _extract_oracle_query(full_sql: str) -> str:
    # Take everything after "2) ORACLE SECTION" and find the first "WITH\nparams AS"
    marker = "2) ORACLE SECTION"
    idx = full_sql.find(marker)
    if idx < 0:
        raise ValueError("Could not find Oracle section marker in SQL file.")
    tail = full_sql[idx:]
    m = re.search(r"(?is)\bWITH\s+params\s+AS\s*\(", tail)
    if not m:
        raise ValueError("Could not locate Oracle query starting at 'WITH params AS (...)'.")
    return tail[m.start() :].strip()


def _oracle_sql_with_bound_po_list(oracle_sql: str, asof_date: Optional[str]) -> str:
    """
    Replace the 'duckdb_po_list AS (...)' CTE with a bind-friendly collection table.
    Also optionally replace params(asof_dt/lookback_dt) with a bound :asof_dt.
    """
    # Replace duckdb_po_list CTE
    oracle_sql = re.sub(
        r"(?is)duckdb_po_list\s+AS\s*\(\s*.*?\)\s*,\s*hdr_candidates\s+AS",
        "duckdb_po_list AS (\n  SELECT column_value AS po_id\n  FROM TABLE(:po_list)\n),\n"
        "hdr_candidates AS",
        oracle_sql,
        count=1,
    )

    if asof_date:
        # Replace params CTE with bind-based version
        oracle_sql = re.sub(
            r"(?is)params\s+AS\s*\(\s*SELECT\s+TRUNC\s*\(\s*SYSDATE\s*\)\s+AS\s+asof_dt\s*,\s*ADD_MONTHS\s*\(\s*TRUNC\s*\(\s*SYSDATE\s*\)\s*,\s*-12\s*\)\s+AS\s+lookback_dt\s*FROM\s+dual\s*\)\s*,",
            "params AS (\n"
            "  SELECT TRUNC(TO_DATE(:asof_dt, 'YYYY-MM-DD')) AS asof_dt,\n"
            "         ADD_MONTHS(TRUNC(TO_DATE(:asof_dt, 'YYYY-MM-DD')), -12) AS lookback_dt\n"
            "  FROM dual\n"
            "),\n",
            oracle_sql,
            count=1,
        )

    return oracle_sql


@dataclass
class DuckDBResult:
    label: str
    columns: List[str]
    rows: List[Sequence[object]]
    runtime_ms: int


def _run_duckdb_statements(con: duckdb.DuckDBPyConnection, stmts: Sequence[str], limit_rows: int = 5000) -> Tuple[List[DuckDBResult], List[str]]:
    results: List[DuckDBResult] = []
    po_list: List[str] = []
    for idx, stmt in enumerate(stmts, 1):
        t0 = time.time()
        cur = con.execute(stmt)
        cols = [d[0] for d in cur.description] if cur.description else []
        rows = cur.fetchall() if cols else []
        runtime_ms = int((time.time() - t0) * 1000)

        # heuristic: PO list statement is a single column named po_id
        if len(cols) == 1 and cols[0].lower() == "po_id" and idx == len(stmts):
            po_list = [str(r[0]) for r in rows if r and r[0] is not None]
            label = f"{idx}_duckdb_po_list"
            rows_written = rows[: min(len(rows), 50)]  # don't spam workbook
            results.append(DuckDBResult(label=label, columns=cols, rows=rows_written, runtime_ms=runtime_ms))
        else:
            label = f"{idx}_duckdb_stmt"
            results.append(
                DuckDBResult(
                    label=label,
                    columns=cols,
                    rows=rows[: min(len(rows), limit_rows)],
                    runtime_ms=runtime_ms,
                )
            )
    return results, po_list


def _oracle_connect(user: str, password: str, connect_str: str):
    if oracledb is None:
        raise RuntimeError("Missing dependency: oracledb. Install with: pip install oracledb")
    return oracledb.connect(user=user, password=password, dsn=connect_str)


def main() -> int:
    ap = argparse.ArgumentParser(description="Run DuckDB + Oracle grand reconciliation from grand_recon.sql")
    ap.add_argument("--sql-file", default=os.path.join("bh_oracle_extracts", "grand_recon.sql"), help="Path to grand_recon.sql")
    ap.add_argument("--duckdb", required=True, help="Path to DuckDB database file (.duckdb)")
    ap.add_argument("--duckdb-read-only", action="store_true", help="Open DuckDB in read-only mode")
    ap.add_argument("--oracle-user", required=True)
    ap.add_argument("--oracle-password", required=True)
    ap.add_argument("--oracle-connect", required=True, help="host:port/service_name")
    ap.add_argument("--asof-date", default=None, help="YYYY-MM-DD (optional; if omitted Oracle uses SYSDATE)")
    ap.add_argument("--output", required=True, help="Output .xlsx path (or a directory for CSV fallback)")
    ap.add_argument("--oracle-chunk-size", type=int, default=2000, help="POs per Oracle query chunk (default: 2000)")
    ap.add_argument("--oracle-max-rows", type=int, default=200000, help="Max per-PO rows written to Excel (default: 200k)")
    args = ap.parse_args()

    sql_path = os.path.abspath(args.sql_file)
    duckdb_path = os.path.abspath(args.duckdb)
    out_path = os.path.abspath(args.output)

    full_sql = open(sql_path, "r", encoding="utf-8").read()
    duckdb_section = _extract_section(full_sql, "1) DUCKDB SECTION", "2) ORACLE SECTION")
    duckdb_stmts = split_sql_statements(duckdb_section)
    if not duckdb_stmts:
        raise SystemExit("No DuckDB statements found in grand_recon.sql")

    oracle_query = _extract_oracle_query(full_sql)
    oracle_query = _oracle_sql_with_bound_po_list(oracle_query, args.asof_date)

    # ---- DuckDB run ----
    dcon = duckdb.connect(duckdb_path, read_only=args.duckdb_read_only)
    duckdb_results, po_list = _run_duckdb_statements(dcon, duckdb_stmts)
    if not po_list:
        raise SystemExit("DuckDB PO list is empty; cannot run Oracle section.")

    # ---- Oracle run ----
    ocon = _oracle_connect(args.oracle_user, args.oracle_password, args.oracle_connect)
    ocur = ocon.cursor()

    oracle_rows: List[Sequence[object]] = []
    oracle_cols: List[str] = []

    try:
        for i in range(0, len(po_list), args.oracle_chunk_size):
            chunk = po_list[i : i + args.oracle_chunk_size]
            # Bind as Oracle collection
            po_var = ocur.arrayvar(oracledb.DB_TYPE_VARCHAR, chunk, type_name="SYS.ODCIVARCHAR2LIST")
            bind = {"po_list": po_var}
            if args.asof_date:
                bind["asof_dt"] = args.asof_date

            t0 = time.time()
            ocur.execute(oracle_query, bind)
            if not oracle_cols:
                oracle_cols = [d[0] for d in ocur.description]
            oracle_rows.extend(ocur.fetchall())
            _ = int((time.time() - t0) * 1000)
    finally:
        try:
            ocur.close()
        except Exception:
            pass
        try:
            ocon.close()
        except Exception:
            pass

    # Sort by total_remaining_amt desc if present
    try:
        idx_total = [c.lower() for c in oracle_cols].index("total_remaining_amt")
        oracle_rows.sort(key=lambda r: (r[idx_total] or 0), reverse=True)
    except Exception:
        pass

    # ---- Write output ----
    if out_path.lower().endswith(".xlsx"):
        if not _HAVE_OPENPYXL:
            raise SystemExit("openpyxl is not installed. Install with: pip install openpyxl  (or pass --output as a directory to write CSVs)")

        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)

        summary = wb.create_sheet("Summary")
        summary.append(["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        summary.append(["SQL file", sql_path])
        summary.append(["DuckDB", duckdb_path])
        summary.append(["Oracle connect", args.oracle_connect])
        summary.append(["As-of date", args.asof_date or "SYSDATE"])
        summary.append([])
        summary.append(["DuckDB PO list size", len(po_list)])
        summary.append(["Oracle rows returned", len(oracle_rows)])

        for r in duckdb_results:
            ws = wb.create_sheet(r.label[:31])
            ws["A1"] = "Runtime (ms)"
            ws["B1"] = r.runtime_ms
            if r.columns:
                _write_table(ws, 3, 1, r.columns, r.rows)
            else:
                ws["A3"] = "(No result set)"

        ws_or = wb.create_sheet("Oracle_PO_Recon")
        ws_or["A1"] = "Rows (written)"
        ws_or["B1"] = min(len(oracle_rows), args.oracle_max_rows)
        if oracle_cols:
            _write_table(ws_or, 3, 1, oracle_cols, oracle_rows[: args.oracle_max_rows])

        wb.save(out_path)
    else:
        # CSV fallback: --output is treated as a directory
        out_dir = out_path
        os.makedirs(out_dir, exist_ok=True)

        _write_csv(
            os.path.join(out_dir, "summary.csv"),
            ["key", "value"],
            [
                ("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                ("sql_file", sql_path),
                ("duckdb", duckdb_path),
                ("oracle_connect", args.oracle_connect),
                ("asof_date", args.asof_date or "SYSDATE"),
                ("duckdb_po_list_size", len(po_list)),
                ("oracle_rows_returned", len(oracle_rows)),
            ],
        )

        for r in duckdb_results:
            _write_csv(os.path.join(out_dir, f"{r.label}.csv"), r.columns, r.rows)

        _write_csv(
            os.path.join(out_dir, "oracle_po_recon.csv"),
            oracle_cols,
            oracle_rows[: args.oracle_max_rows],
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

